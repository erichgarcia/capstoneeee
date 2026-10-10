import json,re,copy
from pathlib import Path
from docx import Document
from docx.shared import Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
p=Path(r'C:/Users/Erich Garcia/Downloads/Capstone_Progress_Final-main (2)/outputs/manuscript-update')
d=Document(p/'source.docx'); original=list(d.paragraphs)
updates=json.loads((p/'updates.json').read_text())
def text(q,value):
 runs=[r for r in q.runs if not r._r.xpath('.//w:drawing')]
 if runs:
  runs[0].text=value
  for r in runs[1:]:r.text=''
 else:q.add_run(value)
for i,value in updates.items():text(original[int(i)],value)
for q in original:
 if re.search(r'Figure\s+(\d+)',q.text):
  updated=re.sub(r'(Figure\s+)(\d+)',lambda m:m[1]+str(int(m[2])+3) if int(m[2])>=39 else m[0],q.text)
  if updated!=q.text:text(q,updated)
def picture(q,name):
 for r in list(q.runs):
  if r._r.xpath('.//w:drawing'):q._p.remove(r._r)
 q.add_run().add_picture(str(p/name),width=Inches(6))
 q.alignment=WD_ALIGN_PARAGRAPH.CENTER
picture(original[529],'Screenshot 2026-10-09 164255.png')
text(original[530],'Figure 20. Registrar Queue Type Selection with Validated Student Information')
# Separate the former combined picture/caption paragraph.
q=d.add_paragraph();original[570]._p.addprevious(q._p)
picture(q,'Screenshot 2026-10-09 164342.png')
for r in list(original[570].runs):
 if r._r.xpath('.//w:drawing'):original[570]._p.remove(r._r)
text(original[570],'Figure 28. Registrar Window Selection')
new=[('Screenshot 2026-10-09 164418.png','Predicted Waiting Time Interface','The Predicted Waiting Time Interface presents estimated waiting times for Cashier, Registrar, ITM, and Admission. Each department card shows the current waiting queue, active service windows, average service time, and traffic status. The estimates use current queue conditions and historical service performance to support operational monitoring. The values shown represent the captured system state and change as queue activity changes.'),('Screenshot 2026-10-09 164549.png','Peak Hour Prediction Interface','The Peak Hour Prediction Interface identifies an upcoming busy period for each of the four departments using historical queue arrivals and recent queue activity. Each card presents the predicted time interval, expected arrival volume, peak-risk level, confidence indicator, and historical pattern. These forecasts help administrators anticipate demand and plan department service coverage. The displayed indicators are estimates rather than guarantees of future arrivals or verified prediction accuracy.'),('Screenshot 2026-10-09 164628.png','Smart Queue Congestion Risk Interface','The Smart Queue Congestion Risk Interface summarizes congestion conditions for Cashier, Registrar, ITM, and Admission. Department cards combine a risk score with the current waiting queue, predicted delay, active windows, arrival rate, service rate, and queue trend. Alerts highlight conditions such as an approaching peak period. This information supports timely operational decisions while reflecting the current queue state; ITM and Admission operate as independent service departments.')]
anchor=original[629]
for num,(name,title,desc) in enumerate(new,39):
 q=d.add_paragraph();anchor._p.addprevious(q._p);q.paragraph_format.page_break_before=True;picture(q,name);q.paragraph_format.keep_with_next=True
 c=d.add_paragraph();anchor._p.addprevious(c._p);c._p.insert(0,copy.deepcopy(original[530]._p.pPr)) if original[530]._p.pPr is not None else None
 c.alignment=WD_ALIGN_PARAGRAPH.CENTER;c.paragraph_format.keep_with_next=True
 c.add_run('Figure '+str(num)+'. ').bold=True;c.add_run(title).italic=True
 b=d.add_paragraph();anchor._p.addprevious(b._p)
 if original[619]._p.pPr is not None:b._p.insert(0,copy.deepcopy(original[619]._p.pPr))
 b.add_run(desc)
# Keep the next section on a fresh page.
anchor.paragraph_format.page_break_before=True
final=p/'CapstoneCHAPTER1-5_Updated_User_Interface.docx'
d.save(final)
print(final)
print('Paragraphs',len(d.paragraphs),'Inline images',len(d.inline_shapes))
print('New figure captions:',[q.text for q in d.paragraphs if re.match(r'Figure (39|40|41|42|43)\.',q.text)])
