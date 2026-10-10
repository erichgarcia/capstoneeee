import json,sys
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
p=sys.argv[1]
doc=Document();sec=doc.sections[0]
sec.page_width=Inches(8.5);sec.page_height=Inches(11.7)
sec.top_margin=sec.bottom_margin=Inches(.5);sec.left_margin=sec.right_margin=Inches(.55)
normal=doc.styles['Normal'];normal.font.name='Arial';normal.font.size=Pt(10);normal.paragraph_format.space_after=Pt(0)
normal.paragraph_format.line_spacing=1
h=doc.add_paragraph('Software Evaluation Total Tally','Title');h.alignment=WD_ALIGN_PARAGRAPH.CENTER
h.runs[0].font.size=Pt(14);h.runs[0].font.color.rgb=RGBColor(0,0,0)
h.paragraph_format.space_after=Pt(4)
a=doc.add_paragraph('40 client records (38 actual + 2 sample clients) and 10 IT experts');a.alignment=WD_ALIGN_PARAGRAPH.CENTER;a.paragraph_format.space_after=Pt(3)
a=doc.add_paragraph('Provisional totals include two synthetic examples. Unresolved original ratings are excluded.');a.alignment=WD_ALIGN_PARAGRAPH.CENTER;a.runs[0].font.size=Pt(9);a.paragraph_format.space_after=Pt(6)
rows=json.load(open(p+'/word_rows.json'));t=doc.add_table(rows=0,cols=6);t.autofit=False;t.alignment=WD_TABLE_ALIGNMENT.CENTER
widths=[3.65,.75,.75,.75,.75,.75]
for c,w in zip(t.columns,widths):c.width=Inches(w)
for i,vals in enumerate(rows):
 cells=t.add_row().cells;section=i>0 and all(v is None for v in vals[1:]);
 for j,(cell,v) in enumerate(zip(cells,vals)):
  cell.width=Inches(widths[j]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
  cell.text='' if v is None or v==0 else str(v)
  par=cell.paragraphs[0];par.alignment=WD_ALIGN_PARAGRAPH.LEFT if section else (WD_ALIGN_PARAGRAPH.RIGHT if j==0 and i>0 else WD_ALIGN_PARAGRAPH.CENTER)
  par.paragraph_format.space_before=Pt(2);par.paragraph_format.space_after=Pt(2)
  for run in par.runs:run.font.size=Pt(10);run.bold=i==0 or j==0
  props=cell._tc.get_or_add_tcPr();b=OxmlElement('w:tcBorders')
  for side in ['top','left','bottom','right']:
   edge=OxmlElement('w:'+side);edge.set(qn('w:val'),'single');edge.set(qn('w:sz'),'4');edge.set(qn('w:color'),'555555');b.append(edge)
  props.append(b)
  if section:
   shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'BFBFBF');props.append(shade)
 rpr=t.rows[-1]._tr.get_or_add_trPr();rpr.append(OxmlElement('w:cantSplit'))
 if i==0:rpr.append(OxmlElement('w:tblHeader'))
a=doc.add_paragraph('Scale follows the reference table: Excellent = 5, Very Good = 4, Good = 3, Fair = 2, Poor = 1. The questionnaire labels ratings 2 and 1 as Poor and Very Poor.');a.paragraph_format.space_before=Pt(5)
for r in a.runs:r.font.size=Pt(8)
doc.save(p+'/evaluation_total_with_2_samples.docx')
